import io
from collections import Counter
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image as PILImage

from .domain import HEADERS, sheet_row, now


def photo_thumbnail(data, size=(500, 340)):
    if not data:
        return None
    im = PILImage.open(io.BytesIO(data)).convert('RGB')
    im.thumbnail(size)
    out = io.BytesIO()
    im.save(out, format='JPEG', quality=72)
    return out.getvalue()


def safe_photo(store, fid):
    try:
        return photo_thumbnail(store.get(fid)) if fid else None
    except Exception:
        return None


def excel(records, store):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.drawing.image import Image
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = '위험성평가'
    cols = HEADERS[:18] + ['개선조치 내용', '처리자', '조치시각', '처리상태', '평가기준', '평가ID']
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='123A63')
        cell.font = Font(name='맑은 고딕', color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 34
    for idx, r in enumerate(records, 2):
        a = r.get('after') or {}
        values = sheet_row(r)[:18] + [a.get('text', ''), a.get('actor', ''), a.get('created_at', ''),
                                     r['status'], r['policy'], r['id']]
        for j, value in enumerate(values, 1):
            if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
                value = "'" + value
            cell = ws.cell(idx, j, value)
            cell.font = Font(name='맑은 고딕', size=10)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if idx % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F1F5FA')
        ws.row_dimensions[idx].height = 100
        for col, fid in [(13, r.get('photo_id')), (18, a.get('photo_id'))]:
            data = safe_photo(store, fid)
            if data:
                pic = Image(io.BytesIO(data))
                ratio = min(150 / pic.width, 110 / pic.height)
                pic.width, pic.height = pic.width * ratio, pic.height * ratio
                ws.add_image(pic, f'{get_column_letter(col)}{idx}')
                ws.cell(idx, col, '')
            else:
                ws.cell(idx, col, '사진 불러오기 실패' if fid else '사진 없음')
    widths = {1:25, 4:24, 6:46, 11:46, 12:38, 13:24, 18:24, 19:46, 20:30, 21:26, 23:28, 24:40}
    for col in range(1, len(cols)+1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 16)
    ws.freeze_panes = 'G2'
    ws.auto_filter.ref = ws.dimensions
    summary = wb.create_sheet('집계')
    summary.append(['생성시각', now()])
    summary.append(['전체 건수', len(records)])
    summary.append(['완료 건수', sum(r['status']=='완료' for r in records)])
    summary.append(['집계 기준', '최초 평가의 당시 등급 / 과거 등급 자동 변경 없음'])
    for grade, count in Counter(r['before']['grade'] for r in records).items():
        summary.append([grade, count])
    summary.column_dimensions['A'].width = 25
    summary.column_dimensions['B'].width = 65
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def pdf(records, store):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
    font_path = next((p for p in [Path(__file__).parent.parent/'assets'/'NanumGothic.ttf',
        Path('/usr/share/fonts/truetype/nanum/NanumGothic.ttf'), Path('C:/Windows/Fonts/malgun.ttf')] if p.exists()), None)
    if not font_path:
        raise ValueError('한글 PDF 글꼴이 없습니다. fonts-nanum 설치 또는 assets/NanumGothic.ttf가 필요합니다.')
    if 'KywaKorean' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('KywaKorean', str(font_path)))
    normal = ParagraphStyle('body', fontName='KywaKorean', fontSize=9, leading=15, wordWrap='CJK', spaceAfter=7)
    title = ParagraphStyle('title', parent=normal, fontSize=24, leading=34, textColor=colors.HexColor('#143C64'), spaceAfter=20)
    heading = ParagraphStyle('heading', parent=normal, fontSize=13, leading=22, spaceBefore=12)
    p = lambda text, style=normal: Paragraph(escape(str(text)).replace('\n', '<br/>'), style)
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=40, bottomMargin=40)
    count = Counter(r['before']['grade'] for r in records)
    story = [p('KYWA 위험성평가\n및 개선조치 보고서', title), p('한국청소년활동진흥원'),
             p(f'생성시각  {now()}'), Spacer(1, 18),
             p(f"평가 {len(records)}건 · 완료 {sum(r['status']=='완료' for r in records)}건 · 미완료 {sum(r['status']!='완료' for r in records)}건", heading),
             p('최초 평가 당시 등급: ' + ' / '.join(f'{g} {n}건' for g, n in count.items())),
             p('본 보고서는 사용자가 확인·저장한 평가와 조치 이력을 기준으로 생성되었습니다. AI 제안의 관련근거는 원문 확인이 필요합니다. 과거 평가의 당시 등급은 보존됩니다.')]
    for index, record in enumerate(records, 1):
        b, a = record['before'], record.get('after') or {}
        story.extend([PageBreak(), p(f"{index:02d}  {record['facility']} · {b['location']}", heading),
                      p(f"{record['department']}  |  {record['status']}  |  {record['created_at']}"),
                      p(f"평가 ID: {record['id']}  /  기준: {record['policy']}"),
                      p(f"{b['category']}  ·  {b['grade']} {b['score']}점 (빈도 {b['p']} × 강도 {b['s']})", heading),
                      p('위험상황  ' + b['scenario']), p('감소대책  ' + b['solution']), p('관련근거(확인 필요)  ' + b['law'])])
        if a:
            story.extend([p(f"개선 후  {a['grade']} {a['score']}점 (빈도 {a['p']} × 강도 {a['s']})", heading),
                          p('실제 조치내용  ' + a.get('text', '기존 기록에 조치내용 없음')),
                          p(f"처리자: {a.get('actor', '기록 없음')} / {a.get('created_at', '')}")])
        photo_cells = []
        for label, fid in [('개선 전 사진', record.get('photo_id')), ('개선 후 사진', a.get('photo_id'))]:
            data = safe_photo(store, fid)
            if data:
                im = PILImage.open(io.BytesIO(data))
                factor = min(245 / im.width, 165 / im.height)
                photo_cells.append([p(label), Image(io.BytesIO(data), width=im.width*factor, height=im.height*factor)])
            else:
                photo_cells.append([p(label), p('사진 불러오기 실패' if fid else '사진 없음')])
        table = Table([photo_cells], colWidths=[260, 260])
        table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('BOX',(0,0),(-1,-1),0.5,colors.HexColor('#DCE3EA'))]))
        story.append(table)
    def footer(canvas, _doc):
        canvas.setFont('KywaKorean', 8)
        canvas.drawString(36, 22, 'KYWA · AI 위험성평가')
        canvas.drawRightString(A4[0]-36, 22, str(_doc.page))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
