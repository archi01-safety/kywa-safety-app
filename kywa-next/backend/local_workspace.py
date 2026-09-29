"""Create a private local test dataset from an XLSX copy. Never calls Google APIs."""
import argparse
import hashlib
import json
import shutil
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .config import ROOT
from .domain import HEADERS
from .migrate import legacy_record
from .services import prepare_image


def workbook_records(source):
    source = Path(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook['설문지 응답 시트1']
        rows = sheet.iter_rows(values_only=True)
        header = list(next(rows))
        while header and header[-1] is None:
            header.pop()
        if header not in (HEADERS[:13], HEADERS[:18]):
            raise ValueError('A:M 또는 A:R 원본 헤더와 다릅니다. 자동 변환을 중단합니다.')
        records = []
        for row_number, raw in enumerate(rows, 2):
            if not any(v is not None and str(v).strip() for v in raw):
                continue
            if any(v is not None for v in raw[len(header):]):
                raise ValueError(f'{row_number}행: 예상 범위 밖 값이 있습니다.')
            values = [v.strftime('%Y-%m-%d %H:%M:%S') if isinstance(v, (datetime,date))
                      else '' if v is None else v for v in raw[:len(header)]]
            try:
                record = legacy_record(values, f'local-xlsx:{source_hash}:{row_number}')
            except (ValueError,TypeError,IndexError) as exc:
                raise ValueError(f'{row_number}행을 변환할 수 없습니다: {exc}') from exc
            record['local_source_row'] = row_number
            # Keep the original references in legacy_values. Never fetch Drive in local mode.
            record['photo_id'] = ''
            if record.get('after'):
                record['after']['photo_id'] = ''
            records.append(record)
        return records, source_hash
    finally:
        workbook.close()


def prepare_workspace(source, photos, output):
    source, photos, output = Path(source).resolve(), Path(photos).resolve(), Path(output).resolve()
    if not output.is_relative_to((ROOT/'local-data').resolve()):
        raise ValueError('시험 저장 공간은 프로젝트 local-data 안에 지정하세요.')
    records, source_hash = workbook_records(source)
    if not photos.is_dir():
        raise ValueError('pic-test 사진 폴더가 없습니다.')
    # mkdir fails if the workspace already exists: never reset the user's later test edits.
    output.mkdir(parents=True, exist_ok=False)
    shutil.copy2(source, output/'source.xlsx')
    (output/'photos').mkdir()
    cache, matched, missing, errors = {}, [], [], []
    for record in records:
        if not record['legacy_values'][12] or record['legacy_values'][12] == '사진 없음':
            continue
        stamp = datetime.fromisoformat(record['created_at']).strftime('%Y%m%d_%H%M%S')
        filename = f"{record['facility']}_{stamp}.jpg"
        path = photos/filename
        if not path.is_file():
            missing.append(record['local_source_row'])
            continue
        try:
            if filename not in cache:
                processed = prepare_image(path.read_bytes())
                file_id = hashlib.sha256(processed).hexdigest()
                (output/'photos'/f'{file_id}.jpg').write_bytes(processed)
                cache[filename] = file_id
            record['photo_id'] = cache[filename]
            matched.append({'source_row':record['local_source_row'], 'filename':filename})
        except Exception as exc:
            errors.append({'source_row':record['local_source_row'], 'error':type(exc).__name__})
    data = {'records':records, 'events':[]}
    report = {'source_sha256':source_hash, 'records':len(records),
              'sample_photos':len(list(photos.glob('*.jpg'))), 'matched_before_photos':matched,
              'unmatched_before_rows':missing, 'photo_errors':errors,
              'after_photos_unmatched':sum(bool(r.get('after') and r['legacy_values'][17]
                   and r['legacy_values'][17] != '사진 없음') for r in records),
              'photo_matching_rule':'원본 코드의 시설명_제출시각.jpg와 정확히 일치할 때만 연결. 개선 후 사진은 임의 연결하지 않음.'}
    (output/'data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    (output/'import-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',required=True)
    parser.add_argument('--photos',default=str(ROOT.parent/'pic-test'))
    parser.add_argument('--output',default=str(ROOT/'local-data/workbook-test'))
    args = parser.parse_args()
    report = prepare_workspace(args.source,args.photos,args.output)
    print(json.dumps({k:v if not isinstance(v,list) else len(v) for k,v in report.items()},ensure_ascii=True))


if __name__ == '__main__':
    main()
