import copy
import io
import json
import unicodedata
import uuid
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from backend import reports
from backend.domain import HEADERS, sheet_row
from backend.organization import load_organization
from backend.repository import SheetsRepository, DriveStore
from backend.submitter_name import display_name, is_masked_name
from backend.tests.test_organizations import configured, connect, submission


@pytest.mark.parametrize('raw,expected', [
    ('홍','홍'), ('김민','김㉤'), ('홍길동','홍㉠㉢'), ('남궁민','남㉠㉤'),
    ('김하늘별','김㉭㉡㉥'), ('홍까따빠싸짜','홍ㄲㄸㅃㅆㅉ'),
    (' \t홍 길\n동　','홍㉠㉢'), (unicodedata.normalize('NFD','홍길동'),'홍㉠㉢'),
    ('홍'+'가'*99,'홍'+'㉠'*99), ('홍㉠㉢','홍㉠㉢'), ('홍ㄲㄸㅃㅆㅉ','홍ㄲㄸㅃㅆㅉ'),
    ('',''), (None,''), ('홍A',''), ('홍ㄱ',''), ('ㅎ',''), ('홍1',''), ('홍!' ,''),
    ('홍\u200b길동',''), ('홍'*101,''),
])
def test_display_masks_legacy_names_without_mutation(raw, expected):
    assert display_name(raw) == expected
    if expected:
        assert is_masked_name(expected)
        assert display_name(expected) == expected


def test_all_initials_and_kywa_config():
    syllables = ''.join(chr(0xac00+i*588) for i in range(19))
    assert display_name('홍'+syllables) == '홍㉠ㄲ㉡㉢ㄸ㉣㉤㉥ㅃ㉦ㅆ㉧㉨ㅉ㉩㉪㉫㉬㉭'
    field = load_organization('kywa').submitter_name
    assert field.mode == 'required' and field.show_in_reports


@pytest.mark.parametrize('value', ['홍길동', '홍㉠동', ' 홍㉠㉢', '홍㉠㉢ ', '홍ㄱㄷ', '홍'+'㉠'*100,
                                  'Hong', '홍\u200b㉠', '㉠㉢', '홍㉠\n', None, ['홍길동'], {'raw':'홍길동'}])
def test_raw_or_invalid_submission_is_rejected_without_echo(tmp_path, caplog, value):
    c, app, _ = connect(configured(tmp_path))
    body = {**submission(c), 'submitter_name':value}
    result = c.post('/api/assessments', json=body)
    assert result.status_code == 422
    assert set(result.json()) == {'detail'} and isinstance(result.json()['detail'], str)
    assert '홍' not in result.text and 'Hong' not in result.text
    assert '홍' not in caplog.text and 'Hong' not in caplog.text
    assert app.state.repo.records() == []
    c.close()


def test_validation_of_extra_fields_and_broken_json_does_not_echo_name(tmp_path, caplog):
    c, _, _ = connect(configured(tmp_path))
    body = {**submission(c), 'submitter_name':'홍㉠㉢', 'raw_name':'홍길동'}
    response = c.post('/api/assessments', json=body)
    assert response.status_code == 422 and '홍' not in response.text
    response = c.post('/api/assessments', content='{"submitter_name":"홍길동",', headers={'Content-Type':'application/json'})
    assert response.status_code == 422 and '홍' not in response.text and '홍' not in caplog.text
    c.close()


def test_optional_blank_submission(tmp_path):
    c, app, _ = connect(configured(tmp_path, 'optional'))
    assert c.post('/api/assessments', json=submission(c)).status_code == 200
    assert app.state.repo.records()[0]['submitter_name'] == ''
    c.close()


@pytest.mark.parametrize('value', ['홍', '김㉤', '홍ㄲㄸㅃㅆㅉ', '홍'+'㉠'*99])
def test_valid_masked_names_are_saved_unchanged(tmp_path, value):
    c, app, _ = connect(configured(tmp_path))
    response = c.post('/api/assessments', json={**submission(c), 'submitter_name':value})
    assert response.status_code == 200
    assert app.state.repo.records()[0]['submitter_name'] == value
    c.close()


def test_sheet_write_and_admin_exports_match_and_ids_stay_independent(tmp_path, monkeypatch):
    cfg = configured(tmp_path)
    c, app, _ = connect(cfg)
    body = {**submission(c), 'submitter_name':'홍㉠㉢'}
    first = c.post('/api/assessments', json=body)
    second = c.post('/api/assessments', json={**body, 'request_id':str(uuid.uuid4())})
    assert first.status_code == second.status_code == 200
    assert first.json()['ids'] != second.json()['ids']
    records = app.state.repo.records()
    writes = []
    class Sheets:
        def get(self, **_):
            return SimpleNamespace(execute=lambda: {'sheets':[{'properties':{'title':cfg.sheet_name,'sheetId':1}}, {'properties':{'title':cfg.history_sheet_name,'sheetId':2}}]})
        def batchUpdate(self, **kwargs):
            writes.append(kwargs['body'])
            return SimpleNamespace(execute=lambda: {})
    repo = SheetsRepository(cfg)
    monkeypatch.setattr(repo, '_values', lambda _: [HEADERS+['제출자 성명']])
    monkeypatch.setattr('backend.repository.google_client', lambda *_: SimpleNamespace(spreadsheets=lambda: Sheets()))
    repo.apply(records, [])
    rows = writes[0]['requests'][0]['appendCells']['rows']
    assert len(rows) == 2
    for row in rows:
        values = row['values']
        assert len(values) == 24
        assert values[23]['userEnteredValue']['stringValue'] == '홍㉠㉢'
        assert json.loads(values[22]['userEnteredValue']['stringValue'])['submitter_name'] == '홍㉠㉢'
    assert '홍길동' not in json.dumps(writes, ensure_ascii=False)
    c.post('/api/auth/demo', data={'role':'admin'})
    assert all(r['submitter_name'] == '홍㉠㉢' for r in c.get('/api/assessments').json()['items'])
    assert 'submitter_name' not in c.get('/api/overview').text
    for format in ['xlsx','pdf']:
        result = c.post('/api/reports', json={'format':format})
        assert result.status_code == 200
        if format == 'xlsx':
            wb = load_workbook(io.BytesIO(result.content))
            assert wb.active.cell(2,25).value == wb.active.cell(3,25).value == '홍㉠㉢'
        else:
            text = ''.join(p.extract_text() for p in PdfReader(io.BytesIO(result.content)).pages)
            assert text.count('홍㉠㉢') == 2 and '홍길동' not in text
    c.close()


def test_legacy_read_and_reports_mask_without_rewriting_original(tmp_path):
    cfg = configured(tmp_path)
    c, app, _ = connect(cfg)
    c.post('/api/assessments', json={**submission(c),'submitter_name':'홍㉠㉢'})
    record = app.state.repo.records()[0]
    record['submitter_name'] = '홍길동'  # Existing historical record, not a new API input.
    missing = copy.deepcopy(record)
    missing['id'] = 'old-no-name'
    missing.pop('submitter_name')
    invalid = copy.deepcopy(record)
    invalid.update(id='old-unsupported', submitter_name='Old Name')
    app.state.repo.apply([record, missing, invalid], [])
    original = app.state.repo.path.read_bytes()
    c.post('/api/auth/demo', data={'role':'admin'})
    items = {r['id']:r for r in c.get('/api/assessments').json()['items']}
    assert items[record['id']]['submitter_name'] == '홍㉠㉢'
    assert items['old-no-name']['submitter_name'] == items['old-unsupported']['submitter_name'] == ''
    wb = load_workbook(io.BytesIO(reports.excel([record,missing,invalid], DriveStore(cfg), cfg.organization)))
    assert wb.active.cell(2,25).value == '홍㉠㉢'
    assert wb.active.cell(3,25).value is None and wb.active.cell(4,25).value is None
    pdf = reports.pdf([record], DriveStore(cfg), cfg.organization)
    text = ''.join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert '홍㉠㉢' in text and '홍길동' not in text
    assert app.state.repo.path.read_bytes() == original
    c.close()


def test_x_only_legacy_name_is_available_for_masked_display(tmp_path, monkeypatch):
    cfg = configured(tmp_path)
    repo = SheetsRepository(cfg)
    row = ['']*24
    row[18], row[22], row[23] = 'old', json.dumps({'id':'old'}), '홍길동'
    monkeypatch.setattr(repo, '_values', lambda _: [HEADERS+['제출자 성명'],row])
    record = repo.records()[0]
    assert display_name(record['submitter_name']) == '홍㉠㉢'
    assert 'submitter_name' not in json.loads(row[22]) and row[23] == '홍길동'
