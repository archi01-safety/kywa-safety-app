import copy
import io
import json
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.config import Settings
from backend.domain import risk, sheet_row
from backend.main import create_app
from backend.organization import load_organization, Organization, PolicyError, record_policy
from backend.repository import DemoRepository, DriveStore, SheetsRepository, RepositoryError
from backend.services import AIService
from backend import reports


def configured(tmp_path, mode='required'):
    org = load_organization('example').model_dump()
    org['submitter_name']['mode'] = mode
    config = tmp_path / 'config'
    config.mkdir(exist_ok=True)
    (config / 'example.json').write_text(json.dumps(org), encoding='utf-8')
    return Settings(_env_file=None, organization_id='example', organization_config_dir=str(config),
                    app_env='demo', app_base_url='http://testserver', demo_data_path=str(tmp_path/'records.json'))


def connect(cfg, **kwargs):
    app = create_app(cfg, **kwargs)
    client = TestClient(app)
    boot = client.get('/api/bootstrap').json()
    client.headers.update({'Origin':cfg.app_base_url, 'X-CSRF-Token':boot['csrf']})
    return client, app, boot


def submission(c):
    response = c.post('/api/analyses', data={'facility':'시험시설', 'department':'시험부서', 'description':'시험 난간'})
    assert response.status_code == 200, response.text
    draft = response.json()
    assert draft['items'][0]['score'] == 6  # Matrix is deliberately NOT p*s (9).
    assert draft['items'][0]['grade'] == '높음'
    return {'request_id':str(uuid.uuid4()), 'draft_id':draft['draft_id'],
            'items':[dict(index=0, location='시험장소', scenario='시험위험', solution='시험조치')]}


def test_kywa_all_cells_match_existing_policy():
    policy = load_organization('kywa').policy
    for p in range(1, 6):
        for s in range(1, 5):
            old = risk(p, s)
            assert {k:v for k,v in policy.risk(p, s).items() if k in old} == old
    for p,s in [(True,1), (6,1), (1,5), ('1',1)]:
        with pytest.raises(PolicyError): policy.risk(p,s)


@pytest.mark.parametrize('change', ['missing-cell', 'duplicate-cell', 'wrong-grade', 'wrong-product'])
def test_invalid_configuration_fails_closed(change):
    org = load_organization('kywa').model_dump()
    cells = org['policy']['cells']
    if change == 'missing-cell': cells.pop()
    if change == 'duplicate-cell': cells.append(cells[0])
    if change == 'wrong-grade': cells[0]['grade'] = '없는 등급'
    if change == 'wrong-product': cells[0]['score'] = 99
    with pytest.raises(ValueError): Organization.model_validate(org)


def test_unknown_id_and_demo_only_refused(tmp_path):
    for identity in ['missing', '../kywa']:
        with pytest.raises(ValueError): Settings(_env_file=None, organization_id=identity).validate_runtime()
    cfg = configured(tmp_path)
    cfg.app_env = 'production'
    with pytest.raises(ValueError, match='시험 기관'): cfg.validate_runtime()


@pytest.mark.parametrize('mode', ['disabled','optional','required'])
def test_name_modes_and_private_exports(tmp_path, mode):
    cfg = configured(tmp_path, mode)
    c, app, boot = connect(cfg)
    assert boot['facilities'] == ['시험시설']
    assert boot['organization']['submitter_name']['mode'] == mode
    assert c.get('/api/overview').json()['total'] == 0
    body = submission(c)
    if mode == 'required':
        assert c.post('/api/assessments', json=body).status_code == 422
    if mode == 'disabled':
        assert c.post('/api/assessments', json={**body, 'submitter_name':'가㉤'}).status_code == 422
    name = '' if mode == 'disabled' else '홍㉠㉢'
    body['submitter_name'] = name
    assert c.post('/api/assessments', json=body).status_code == 200
    assert c.post('/api/assessments', json=body).json()['replayed']
    record = app.state.repo.records()[0]
    assert record['organization_id'] == 'example'
    assert record['policy_snapshot']['method'] == 'matrix'
    assert record['submitter_name'] == name
    assert sheet_row(record, True)[23] == name
    assert name == '' or name not in c.get('/api/overview').text
    assert c.get('/api/assessments').status_code == 401
    assert c.post('/api/auth/demo', data={'role':'admin'}).status_code == 200
    shown = c.get('/api/assessments').json()['items'][0]
    assert shown.get('submitter_name', '') == name
    wb = load_workbook(io.BytesIO(reports.excel([record], DriveStore(cfg), cfg.organization)))
    assert ('제출자 성명' in [x.value for x in wb.active[1]]) == (mode != 'disabled')
    if name: assert wb.active.cell(2,25).value == name
    cfg.organization.submitter_name.show_in_reports = False
    hidden = load_workbook(io.BytesIO(reports.excel([record], DriveStore(cfg), cfg.organization)))
    assert '제출자 성명' not in [x.value for x in hidden.active[1]]
    c.close()


def test_snapshot_preserved_after_policy_change_and_range_checked(tmp_path):
    cfg = configured(tmp_path)
    c, app, _ = connect(cfg)
    body = {**submission(c), 'submitter_name':'가㉤'}
    rid = c.post('/api/assessments', json=body).json()['ids'][0]
    # New deployment settings must not change the previously recorded matrix.
    cfg.organization.policy.version = 'EXAMPLE-v2'
    for cell in cfg.organization.policy.cells: cell.score = 99
    c.post('/api/auth/demo', data={'role':'admin'})
    response = c.post(f'/api/assessments/{rid}/reassess', data={'text':'시험 조치'})
    result = response.json()
    assert result['score'] == 4
    action = dict(request_id=str(uuid.uuid4()), draft_id=result['draft_id'], revision=result['revision'], p=4, s=1)
    assert c.post(f'/api/assessments/{rid}/actions', json=action).status_code == 422
    action.update(p=1,s=2)
    saved = c.post(f'/api/assessments/{rid}/actions', json=action)
    assert saved.status_code == 200, saved.text
    record = app.state.repo.records()[0]
    assert record['before']['score'] == 6 and record['after']['score'] == 3
    assert record['after']['policy'] == 'EXAMPLE-3X3-v1'
    assert record['after']['tone'] == 1
    with pytest.raises(PolicyError): record_policy({'policy':'UNKNOWN'}, cfg.organization)
    c.close()


def test_prompt_uses_org_policy_and_excludes_submitter(tmp_path, monkeypatch):
    cfg = configured(tmp_path)
    cfg.app_env = 'development'  # Provider is stubbed; no external call.
    ai = AIService(cfg)
    captured = []
    def generate(prompt, schema, photo):
        captured.append(prompt)
        value = dict(p=3, s=3, rationale='시험') if len(captured)>1 else {'items':[dict(category='시험 분류',location='위치',scenario='상황',solution='조치',p=3,s=3)]}
        return schema.model_validate(value)
    monkeypatch.setattr(ai, '_generate', generate)
    ai.analyze('시험시설','시험부서','시험 설명')
    ai.reassess({'policy':cfg.organization.policy.version, 'before':{'scenario':'시험'}, 'submitter_name':'전송금지성명'},'조치')
    assert all('EXAMPLE-3X3-v1' in p and '전송금지성명' not in p for p in captured)
    assert '한국청소년활동진흥원' not in captured[0]


def test_name_sheet_schema_and_foreign_org_guard(tmp_path, monkeypatch):
    from backend.domain import HEADERS
    cfg = configured(tmp_path)
    repo = SheetsRepository(cfg)
    monkeypatch.setattr(repo, '_values', lambda _: [HEADERS])
    from backend.sheet_schema import Layout, SheetSchemaError
    assert repo.records() == []  # Existing data can be read before preparing new columns.
    with pytest.raises(SheetSchemaError, match='X1'): Layout(HEADERS,cfg.organization).require_writable()
    record = {'id':'test','organization_id':'kywa'}
    row = ['']*24
    row[18],row[22] = 'test', json.dumps(record)
    monkeypatch.setattr(repo, '_values', lambda _: [HEADERS+['제출자 성명'],row])
    with pytest.raises(SheetSchemaError, match='다른 기관'): repo.records()


def test_name_column_migration_preserves_existing_columns():
    from backend.domain import HEADERS
    from backend.migrate import migration_rows
    row = ['']*24
    row[18], row[22] = 'id1', json.dumps({'id':'id1'})
    assert migration_rows([HEADERS, row], 'test', include_name=True) == []
    row[23] = 'existing content'
    with pytest.raises(ValueError, match='X열'):
        migration_rows([HEADERS, row], 'test', include_name=True)
    assert migration_rows([HEADERS+['제출자 성명'], row], 'test', include_name=True) == []


def test_pdf_name_visibility_and_policy(tmp_path):
    from pypdf import PdfReader
    cfg = configured(tmp_path)
    c, app, _ = connect(cfg)
    c.post('/api/assessments', json={**submission(c), 'submitter_name':'홍㉠㉢'})
    records = app.state.repo.records()
    pdf = reports.pdf(records, DriveStore(cfg), cfg.organization)
    text = ''.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
    assert '홍㉠㉢' in text and 'EXAMPLE-3X3-v1' in text
    cfg.organization.submitter_name.show_in_reports = False
    pdf = reports.pdf(records, DriveStore(cfg), cfg.organization)
    text = ''.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
    assert '홍㉠㉢' not in text
    c.close()


def test_historical_kywa_without_snapshot_uses_frozen_v1():
    org = load_organization('kywa')
    org.policy.version = 'KYWA-v2'
    for cell in org.policy.cells: cell.score = 999
    old = record_policy({'policy':'KYWA-2026-6LEVEL-v1'}, org)
    assert old.risk(3, 4)['score'] == 12
