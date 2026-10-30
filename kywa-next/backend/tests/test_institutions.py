import copy
import io
import json
import uuid
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.branding import palette, contrast, mix
from backend.check_connections import INSTITUTIONS, inspect, probe, target_settings, header_edit_allowed
from backend.config import Settings
from backend.domain import HEADERS, HISTORY_HEADERS
from backend.migrate import preparation_plan, prepare_storage
from backend.organization import load_organization, record_policy
from backend.repository import SheetsRepository, RepositoryError, DriveStore
from backend.sheet_schema import Layout, LABELS, MANAGED, base_keys, expanded_header, SheetSchemaError
from backend.tests.test_organizations import connect


ACCOUNT='kywa-safety-check@gen-lang-client-0053629470.iam.gserviceaccount.com'


def settings(identity='khidi'):
    cfg=Settings(_env_file=None,organization_id=identity,app_base_url='http://testserver',
        google_service_account_json=json.dumps(dict(type='service_account',client_email=ACCOUNT,private_key='test-only')))
    cfg.validate_runtime()
    return cfg


def old_row(org):
    values=['2026. 5. 12 오후 3:24:15',org.facilities[0],org.departments[0],
            '복도','보행 안전','미끄러움',5,3,15,'보통','바닥 정리','현장 확인','사진 없음']
    if org.storage.legacy_name_column: values.insert(3,'홍길동')
    return values


def current_record(org):
    return dict(id=str(uuid.uuid4()),organization_id=org.id,created_at='2026-10-29T12:00:00+09:00',
        facility=org.facilities[0],department=org.departments[0],submitter_name='홍㉠㉢',
        before=dict(location='본관',category='보행 안전',scenario='단차',solution='보수',law='확인',**org.policy.risk(5,3)),
        after=None,photo_id='',policy=org.policy.version,policy_snapshot=org.policy.model_dump(),status='접수',revision=1)


class MemorySheets:
    """Models appendCells, targeted updates and old-client append ordering, not Google permissions."""
    def __init__(self,cfg,rows):
        self.cfg=cfg
        self.data={cfg.sheet_name:copy.deepcopy(rows),cfg.history_sheet_name:[HISTORY_HEADERS.copy()]}
        self.ids={cfg.sheet_name:1,cfg.history_sheet_name:2}
        self.requests=[]
        self.before_batch=None

    def spreadsheets(self): return self
    def values(self): return self

    def get(self,**kwargs):
        if 'range' in kwargs:
            name=kwargs['range'].split('!')[0][1:-1].replace("''", "'")
            return SimpleNamespace(execute=lambda:dict(values=copy.deepcopy(self.data[name])))
        return SimpleNamespace(execute=lambda:dict(sheets=[{'properties':{'title':name,'sheetId':sid,
            'gridProperties':{'columnCount':26}}} for name,sid in self.ids.items()]))

    def batchUpdate(self,**kwargs):
        def execute():
            if self.before_batch:
                callback,self.before_batch=self.before_batch,None
                callback()
            requests=kwargs['body']['requests']
            self.requests.extend(copy.deepcopy(requests))
            replies=[]
            for req in requests:
                if 'appendDimension' in req: continue
                if 'addSheet' in req:
                    prop=req['addSheet']['properties']
                    sid=prop.get('sheetId',max(self.ids.values())+1)
                    self.ids[prop['title']]=sid
                    self.data[prop['title']]=[]
                    replies.append({'addSheet':{'properties':{'sheetId':sid}}})
                    continue
                if 'deleteSheet' in req:
                    name=next(n for n,sid in self.ids.items() if sid==req['deleteSheet']['sheetId'])
                    del self.ids[name],self.data[name]
                    continue
                item=req.get('updateCells') or req['appendCells']
                start=item.get('start',{})
                name=next(n for n,sid in self.ids.items() if sid==start.get('sheetId',item.get('sheetId')))
                values=[[next(iter(c['userEnteredValue'].values())) for c in r['values']] for r in item['rows']]
                if 'appendCells' in req:
                    self.data[name].extend(values)
                else:
                    for i,row in enumerate(values,start['rowIndex']):
                        while len(self.data[name])<=i: self.data[name].append([])
                        col=start['columnIndex']
                        self.data[name][i].extend(['']*max(0,col+len(row)-len(self.data[name][i])))
                        self.data[name][i][col:col+len(row)]=row
            return {'replies':replies}
        return SimpleNamespace(execute=execute)


def wire(monkeypatch,cfg,rows):
    memory=MemorySheets(cfg,rows)
    monkeypatch.setattr('backend.repository.google_client',lambda *_:memory)
    monkeypatch.setattr('backend.migrate.google_client',lambda *_:memory)
    return memory,SheetsRepository(cfg)


@pytest.mark.parametrize('identity',INSTITUTIONS)
def test_all_policies_names_palettes_and_public_bootstrap(identity,tmp_path):
    cfg=settings(identity)
    cfg.demo_data_path=str(tmp_path/'data.json')
    org=cfg.organization
    assert org.submitter_name.mode=='required' and org.submitter_name.show_in_reports
    for p in range(1,6):
        for s in range(1,5):
            value=org.policy.risk(p,s)
            assert value['score']==p*s
            bounds=([(2,'매우 낮음'),(4,'낮음'),(8,'보통'),(12,'약간 높음'),(15,'높음'),(20,'매우 높음')] if identity=='fowi' else
                    [(5,'매우 낮음'),(11,'낮음'),(15,'보통'),(18,'높음'),(20,'매우 높음')] if identity=='khidi' else
                    [(3,'매우 낮음'),(6,'낮음'),(8,'보통'),(12,'약간 높음'),(15,'높음'),(20,'매우 높음')] if identity=='kywa' else
                    [(3,'매우 낮음'),(6,'낮음'),(8,'보통'),(15,'높음'),(20,'매우 높음')])
            assert value['grade']==next(g for high,g in bounds if p*s<=high)
    c,_,boot=connect(cfg)
    public=boot['organization']
    assert 'storage' not in public and 'source' not in public
    assert cfg.spreadsheet_id not in json.dumps(boot) and ACCOUNT not in json.dumps(boot)
    colors=public['branding']['colors']
    assert colors==palette(org.branding.primary)
    for state in ('button','hover'): assert contrast(colors[state],colors['buttonText'])>=4.5
    for bg in ('#181B20','#242830','#2C313B',mix(colors['original'],'#242830',.88)):
        assert contrast(colors['darkText'],bg)>=4.5
    for bg in ('#FFFFFF','#F4F4F2','#F0F1F3',mix(colors['original'],'#FFFFFF',.92)):
        assert contrast(colors['lightText'],bg)>=4.5
    assert public['branding']['logo']==('/kywa_logo.png' if identity=='kywa' else '')
    c.close()


def test_shared_folders_and_deduplicated_departments():
    assert load_organization('cayf').storage.drive_folder_id==load_organization('knf').storage.drive_folder_id
    assert load_organization('kohi').storage.drive_folder_id==load_organization('khidi').storage.drive_folder_id
    assert load_organization('kohi').departments.count('운영지원부')==1


@pytest.mark.parametrize('identity',INSTITUTIONS)
def test_mixed_rows_original_order_unknown_columns_and_stable_ids(identity,monkeypatch):
    cfg=settings(identity)
    org=cfg.organization
    header=[LABELS[k] for k in base_keys(org)]
    header[2]='담당 팀명'
    if identity=='khidi': header[3]='이름'  # Confirmed existing D1 header.
    header.append('기존 메모')
    row=old_row(org)+['이 값 보존']
    rows=[header,row,row.copy()]
    expanded=expanded_header(rows,org)
    assert expanded[:len(header)]==header
    layout=Layout(expanded,org)
    assert layout.columns['submitter_name']==(3 if identity=='khidi' else 24)
    record=current_record(org)
    memory,repo=wire(monkeypatch,cfg,[expanded,row,row.copy(),layout.new_row(record)])
    records=repo.records()
    assert len({r['id'] for r in records})==3
    assert [r['id'] for r in repo.records()]==[r['id'] for r in records]
    old=records[0]
    assert old['before']['score']==15 and old['before']['grade']=='보통'
    assert record_policy(old,org).version==org.policy.version
    old.update(status='조치 중',revision=2,after=dict(text='시험 보수',policy=org.policy.version,**org.policy.risk(1,1)))
    # An old app appends between new-app read and batch execution.
    memory.before_batch=lambda:memory.data[cfg.sheet_name].append(row.copy())
    repo.apply([old],[])
    assert memory.data[cfg.sheet_name][1][:len(row)]==row
    assert memory.data[cfg.sheet_name][-1]==row
    updates=[r['updateCells'] for r in memory.requests]
    assert {x['start']['columnIndex'] for x in updates}=={layout.columns[k] for k in MANAGED}
    reread=repo.records()
    assert reread[0]['id']==old['id'] and reread[0]['revision']==2
    assert reread[0]['before']['grade']=='보통'
    if identity=='khidi':
        assert memory.data[cfg.sheet_name][1][3]=='홍길동'
        assert reread[0]['submitter_name']=='홍㉠㉢'
        assert '홍길동' not in json.dumps(updates,ensure_ascii=False)


@pytest.mark.parametrize('change',['source','move','delete','managed'])
def test_stale_legacy_write_conflicts(monkeypatch,change):
    cfg=settings()
    row=old_row(cfg.organization)
    header=expanded_header([[LABELS[k] for k in base_keys(cfg.organization)],row],cfg.organization)
    mem,repo=wire(monkeypatch,cfg,[header,row])
    record=repo.records()[0]
    record.update(status='조치 중',revision=2)
    if change=='source': mem.data[cfg.sheet_name][1][4]='수정된 원본'
    if change=='move': mem.data[cfg.sheet_name].insert(1,[])
    if change=='delete': mem.data[cfg.sheet_name].pop()
    if change=='managed':
        other=repo.records()[0]
        other.update(status='조치 중',revision=2)
        repo.apply([other],[])
    count=len(mem.requests)
    with pytest.raises(RepositoryError,match='변경'): repo.apply([record],[])
    assert len(mem.requests)==count


def test_schema_conflicts_are_not_silently_overwritten():
    org=load_organization('kywa')
    with pytest.raises(SheetSchemaError): expanded_header([HEADERS,['']*23+['기존데이터']],org)
    with pytest.raises(SheetSchemaError): Layout(HEADERS+['성명','제출자 성명'],org)
    with pytest.raises(SheetSchemaError): Layout(HEADERS[:2]+['성명']+HEADERS[2:],org)
    with pytest.raises(ValueError): preparation_plan([HEADERS],[['다른 용도']],org)


def test_protected_header_requires_owner_action():
    sheet={'protectedRanges':[{'range':{'startRowIndex':0,'endRowIndex':1}}]}
    assert not header_edit_allowed(sheet,18,24)
    assert header_edit_allowed(sheet,24,24)  # Nothing to extend.
    sheet['protectedRanges'][0]['warningOnly']=True
    assert header_edit_allowed(sheet,18,24)
    sheet['protectedRanges'][0].update(warningOnly=False,requestingUserCanEdit=True)
    assert header_edit_allowed(sheet,18,24)
    sheet={'protectedRanges':[{'range':{'startRowIndex':1,'endRowIndex':100}}]}
    assert header_edit_allowed(sheet,18,24)


def test_header_preparation_is_append_only_and_idempotent(monkeypatch,tmp_path):
    cfg=settings()
    org=cfg.organization
    row=old_row(org)
    header=[LABELS[k] for k in base_keys(org)]
    mem,_=wire(monkeypatch,cfg,[header,row])
    monkeypatch.setattr('backend.migrate.ROOT',tmp_path)
    assert prepare_storage(cfg)['legacy_rows']==1
    assert mem.requests==[]
    result=prepare_storage(cfg,True)
    assert result['applied'] and result['backup']
    assert mem.data[cfg.sheet_name][1]==row
    assert mem.requests[0]['updateCells']['start']['columnIndex']==14
    assert mem.data[cfg.sheet_name][0][:14]==header
    before=len(mem.requests)
    assert prepare_storage(cfg,True)['added_headers']==[]
    assert len(mem.requests)==before


def test_api_legacy_and_new_app_submit_retry_action_exports(monkeypatch,tmp_path):
    cfg=settings()
    cfg.demo_data_path=str(tmp_path/'data.json')
    org=cfg.organization
    row=old_row(org)
    layout=Layout(expanded_header([[LABELS[k] for k in base_keys(org)],row],org),org)
    mem,repo=wire(monkeypatch,cfg,[layout.header,row])
    c,app,_=connect(cfg,repo=repo)
    draft=c.post('/api/analyses',data={'facility':org.facilities[0],'department':org.departments[0],'description':'시험 위험'}).json()
    body=dict(request_id=str(uuid.uuid4()),draft_id=draft['draft_id'],submitter_name='김㉤',
        items=[dict(index=0,location='시험',scenario='위험',solution='조치')])
    mem.before_batch=lambda:mem.data[cfg.sheet_name].append(row.copy())
    new=c.post('/api/assessments',json=body)
    assert new.status_code==200,new.text
    assert c.post('/api/assessments',json=body).json()['replayed']
    assert len(repo.records())==3
    assert c.get('/api/overview').json()['total']==3
    assert 'submitter_name' not in c.get('/api/overview').text
    c.post('/api/auth/demo',data={'role':'admin'})
    listing=c.get('/api/assessments')
    assert '홍길동' not in listing.text and '_sheet_' not in listing.text
    old=next(r for r in listing.json()['items'] if r['policy']=='LEGACY-UNKNOWN')
    assert old['current_risk']['grade']=='보통' and old['submitter_name']=='홍㉠㉢'
    before_id=old['id']
    after=c.post(f'/api/assessments/{before_id}/reassess',data={'text':'시험 보수'}).json()
    action=dict(request_id=str(uuid.uuid4()),draft_id=after['draft_id'],revision=after['revision'],p=1,s=1)
    saved=c.post(f'/api/assessments/{before_id}/actions',json=action)
    assert saved.status_code==200,saved.text
    assert saved.json()['item']['after']['policy']==org.policy.version
    assert c.post(f'/api/assessments/{before_id}/actions',json=action).json()['replayed']
    assert mem.data[cfg.sheet_name][1][:14]==row
    xlsx=c.post('/api/reports',json={'format':'xlsx'})
    assert xlsx.status_code==200,xlsx.text
    ws=load_workbook(io.BytesIO(xlsx.content)).active
    assert {ws.cell(i,25).value for i in range(2,5)}=={'홍㉠㉢','김㉤'}
    pdf=c.post('/api/reports',json={'format':'pdf'})
    assert pdf.status_code==200
    text=''.join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages)
    assert '홍길동' not in text and '홍㉠㉢' in text and '김㉤' in text
    c.close()


def test_drive_same_operation_is_scoped_by_institution(monkeypatch):
    items=[]
    class Drive:
        def files(self): return self
        def list(self,**kwargs):
            q=kwargs['q']
            return SimpleNamespace(execute=lambda:dict(files=[{'id':x['id']} for x in items
                if "value='"+x['appProperties']['safetyOrganization']+"'" in q and "value='"+x['appProperties']['kywaOperation']+"'" in q]))
        def create(self,**kwargs):
            item={**kwargs['body'],'id':str(len(items)+1)}
            items.append(item)
            return SimpleNamespace(execute=lambda:{'id':item['id']})
    monkeypatch.setattr('backend.repository.google_client',lambda *_:Drive())
    operation=str(uuid.uuid4())
    stores=[]
    for identity in ('knf','cayf'):
        cfg=settings(identity)
        cfg.app_env='production'  # Storage API is mocked.
        stores.append(DriveStore(cfg))
    first=stores[0].put(b'test-image',operation)
    assert stores[0].put(b'test-image',operation)==first
    assert stores[1].put(b'test-image',operation)!=first
    assert len(items)==2
    assert items[0]['parents']==items[1]['parents']
    assert items[0]['name'].startswith('knf_') and items[1]['name'].startswith('cayf_')


@pytest.mark.parametrize('value',['{}','not-json',json.dumps(dict(type='service_account',private_key='secret',client_email='wrong@example.org'))])
def test_bad_credentials_fail_safely(value):
    cfg=settings()
    cfg.google_service_account_json=value
    with pytest.raises(ValueError) as exc: cfg.validate_google_identity()
    assert 'secret' not in str(exc.value) and 'wrong@example.org' not in str(exc.value)


def test_production_wrong_org_destination_is_rejected():
    cfg=settings('knf')
    cfg.app_env='production'
    cfg.app_base_url='https://example.org'
    cfg.session_secret='test-only-'*5
    cfg.google_client_id=cfg.google_client_secret=cfg.gemini_api_key=cfg.kosha_api_key='test-only'
    cfg.access_accounts_json=json.dumps({'admin@example.org':{'role':'admin','facilities':[]}})
    cfg.allow_legacy_writes=True
    cfg.spreadsheet_id=load_organization('cayf').storage.spreadsheet_id
    with pytest.raises(ValueError,match='대상'): cfg.validate_runtime()
